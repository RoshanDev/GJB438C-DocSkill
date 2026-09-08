"""Synthetic regressions: a long appendix never satisfies main-body targets."""
import json
import os
from pathlib import Path
import shutil
import yaml
from zipfile import ZipFile

from docx import Document
from lxml import etree
import pytest

from gjb438c_suite import cli, suite, volume
from gjb438c_suite.audit_docx import audit_docx
from gjb438c_suite.content_scope import (
    PAGE_COUNT_SCOPE, MAIN_BOOKMARK, BACK_BOOKMARK, split_content, headings,
)
from gjb438c_suite.finalize import refresh_toc_cache
from gjb438c_suite.import_word import import_word
from gjb438c_suite.markdown_doc import parse_markdown, split_front_matter
from gjb438c_suite.profile_quality import audit_profile_document
from gjb438c_suite.registry import iter_document_types
from gjb438c_suite.render import render_document
from test_review_round_three import _passing_suite

SOURCE = Path(__file__).resolve().parents[1] / 'examples/SRS.example.md'


def source_file(tmp_path, body):
    path = tmp_path / 'source.md'
    metadata, _, _, errors = split_front_matter(SOURCE.read_text(encoding='utf-8'))
    assert not errors
    metadata['document']['status'] = 'draft'
    metadata['quality'] = {'tier': 'large'}
    path.write_text('---\n' + yaml.safe_dump(metadata, allow_unicode=True, sort_keys=False) + '---\n' + body,
                    encoding='utf-8')
    return path


@pytest.mark.parametrize('title', [
    '# 附录A 数据块', '## 附录 A 数据块', '# **附录 A** 补充', '# 8 附录 B',
    '# 附件1 日志', '# Appendix A', '# APPENDICES', '# Annex 1',
    '# Annexes', '# 8 结构化工程证据', '# 质量门禁数据块',
    '附录 A\n======', '附录 A\n------', '   ## 附录A 数据块',
])
def test_back_matter_never_reopens_as_main(title):
    scope = split_content('# 1 范围\n正文甲。\n' + title + '\n附录内容。\n# 3 需求\n更多附录。')
    assert '正文甲' in scope.main
    assert '附录内容' not in scope.main
    assert '# 3 需求' in scope.back_matter


@pytest.mark.parametrize('fake', [
    '```text\n# 附录 A\n```', '~~~~text\n# 附录 A\n~~~~',
    '<!--\n# 附录 A\n-->', '本章参见附录 A，不是附录标题。',
])
def test_code_comment_or_reference_is_not_a_boundary(fake):
    text = '# 1 范围\n' + fake + '\n正文仍然存在。'
    assert split_content(text).main == text
    assert split_content(text).back_matter == ''


def test_appendix_prose_tables_images_and_evidence_do_not_buy_volume(tmp_path):
    original = '# 1 范围\n这里是简短的真实正文。\n'
    before = parse_markdown(source_file(tmp_path, original))
    appendix = '# 附录A 支撑资料\n' + ('大量补充材料不应计入正文。\n\n' * 500)
    appendix += '\n| 字段 | 值 |\n| --- | --- |\n| X | Y |\n![图](diagram.png)\n'
    appendix += '```gjb-requirement\nid: ANNEX-001\nstatement: 只在附录。\n```\n'
    after = parse_markdown(source_file(tmp_path, original + appendix))
    assert volume.effective_units(after) == volume.effective_units(before)
    issues = volume.markdown_volume_issues(after, 'SRS', 'large', 'review')
    assert any(i['code'] == 'VOLUME_VISIBLE_CONTENT_TOO_THIN' and i['severity'] == 'ERROR' for i in issues)
    assert any(i['code'] == 'VOLUME_BACK_MATTER_EXCLUDED' for i in issues)


def test_prelude_is_not_main_volume(tmp_path):
    body = '# 1 范围\n真实正文。\n'
    source = parse_markdown(source_file(tmp_path, '# 前言\n' + '前言内容' * 1000 + '\n' + body))
    assert volume.visible_markdown_characters(source) < 20


@pytest.mark.parametrize('code', [i.code for i in iter_document_types()])
def test_appendix_cannot_supply_a_missing_required_heading(tmp_path, code):
    source = tmp_path / (code + '.md')
    assert cli.main(['init', '--type', code, '--output', str(source)]) == 0
    text = source.read_text(encoding='utf-8')
    _, offset, _, _, _ = next(h for h in headings(text) if h[2] == 1 and h[3].startswith('2 '))
    text = text[:offset] + '# 附录 A\n' + text[offset:]
    source.write_text(text, encoding='utf-8')
    report = audit_profile_document(source, audit_profile='review')
    assert report.heading_coverage_percent < 100
    assert any(i.code == 'PROFILE_HEADING_MISSING' for i in report.issues)


def test_inline_evidence_is_visible_once_and_not_appended_again(tmp_path):
    body = '# 1 范围\n\n## 1.1 条款\n\n```gjb-requirement\nid: INLINE-001\nstatement: 独立的需求说明。\n```\n\n## 1.2 后续\n正文。\n'
    source = source_file(tmp_path, body)
    word = tmp_path / 'inline.docx'
    render_document(source, word, profile='draft')
    doc = Document(word)
    assert not any('附录 结构化工程证据' in p.text for p in doc.paragraphs)
    assert sum(c.text == 'INLINE-001' for t in doc.tables for row in t.rows for c in row.cells) == 1
    root = doc.element
    flat = ''.join(root.xpath('.//w:t/text()'))
    assert flat.index('1.1 条款') < flat.index('INLINE-001') < flat.index('1.2 后续')
    returned = tmp_path / 'back.md'
    assert import_word(word, returned).exact_round_trip
    assert returned.read_bytes() == source.read_bytes()


def test_nested_scope_bookmarks_do_not_end_style_audit_early(tmp_path):
    word = tmp_path / 'bad-style.docx'
    render_document(SOURCE, word, profile='review')
    doc = Document(word)
    next(p for p in doc.paragraphs if p.text == '3.1 任务管理需求').style = 'Normal'
    doc.save(word)
    report = audit_docx(word, profile='review')
    assert any(i.code == 'BODY_STYLE' for i in report.errors)


def test_old_volume_report_scope_is_rejected(tmp_path, monkeypatch):
    manifest, _ = _passing_suite(tmp_path, monkeypatch)
    report = tmp_path / 'OCD.json'
    data = json.loads(report.read_text(encoding='utf-8'))
    data.pop('page_count_scope')
    report.write_text(json.dumps(data), encoding='utf-8')
    result = suite.audit_suite_manifest(manifest)
    assert not result.passed
    assert any(i.code == 'SUITE_REPORT_MISMATCH' for i in result.issues)


def test_automatic_expansion_is_an_error_not_a_workaround(tmp_path):
    source = parse_markdown(source_file(tmp_path, '# 1 范围\n文档。\n# 附录 展开\n### REQ-001 正文展开 1\n重复字段。'))
    for mode, severity in [('draft', 'WARN'), ('review', 'ERROR'), ('release', 'ERROR')]:
        issues = volume.markdown_volume_issues(source, 'SRS', 'large', mode)
        assert any(i['code'] == 'VOLUME_GENERATED_EXPANSION' and i['severity'] == severity for i in issues)


OFFICE = bool(os.environ.get('GJB_OFFICE_TESTS')) and bool(shutil.which('libreoffice') or shutil.which('soffice'))


@pytest.mark.skipif(not OFFICE, reason='actual Office pagination requested separately')
def test_real_office_short_main_long_appendix_cannot_pass(tmp_path):
    # One paragraph main, hundreds of appendix paragraphs: use real Word/PDF,
    # not mocked page totals. This is test data, never project evidence.
    body = '# 1 范围\n主文档的唯一简短说明。\n\n# 附录 A 补充材料\n\n'
    body += '\n\n'.join(f'附录测试数据 {i}，不属于主文档正文。' * 12 for i in range(100))
    source = source_file(tmp_path, body)
    word = tmp_path / 'long-appendix.docx'
    render_document(source, word, profile='draft')
    refresh_toc_cache(word)
    metrics = volume.rendered_page_metrics(word)
    assert metrics.page_count_scope == PAGE_COUNT_SCOPE
    assert metrics.body_pages == 1
    assert metrics.appendix_pages > 10
    assert metrics.total_pages == metrics.body_start_page - 1 + metrics.body_pages + metrics.appendix_pages
    result = volume.audit_rendered_volume(source, 'SRS', word, tier='large')
    assert not result.passed and result.body_pages == 1
    assert result.appendix_pages > 10
    with ZipFile(word) as z:
        root = etree.fromstring(z.read('word/document.xml'))
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    for name in (MAIN_BOOKMARK, BACK_BOOKMARK):
        assert len(root.xpath('.//w:bookmarkStart[@w:name=$name]', namespaces=ns, name=name)) == 1
    returned = tmp_path / 'returned.md'
    assert import_word(word, returned).exact_round_trip
    assert source.read_bytes() == returned.read_bytes()


@pytest.mark.parametrize('fence', ['```', '~~~~', '````'])
def test_fenced_machine_data_not_narrative(tmp_path, fence):
    base = '# 1 范围\n简短正文。\n'
    before = parse_markdown(source_file(tmp_path, base))
    extra = '\n' + fence + 'text\n' + '字段占位非正文。' * 1000 + '\n' + fence + '\n'
    after = parse_markdown(source_file(tmp_path, base + extra))
    assert volume.visible_markdown_characters(before) == volume.visible_markdown_characters(after)


def test_missing_or_relocated_accounting_marker_fails_closed(tmp_path):
    from docx.oxml.ns import qn
    source = source_file(tmp_path, '# 1 范围\n正文。\n# 附录 A\n补充材料。\n# 附录 B\n其他补充。\n')
    word = tmp_path / 'scopes.docx'
    render_document(source, word, profile='draft')
    class PDF:
        outline = [{'/Title': '1 范围'}, {'/Title': '附录 A'}, {'/Title': '附录 B'}]
        pages = [None] * 100
        def get_destination_page_number(self, item):
            return {'1 范围': 4, '附录 A': 5, '附录 B': 50}[item['/Title']]
    assert volume._page_bounds(PDF(), word) == (4, 5)
    doc = Document(word)
    mark = doc.element.xpath('.//w:bookmarkStart[@w:name="GJB_BACK_MATTER"]')[0]
    mark.getparent().remove(mark)
    next(p for p in doc.paragraphs if p.text == '附录 B')._p.append(mark)
    doc.save(word)
    with pytest.raises(volume.VolumeError, match='第一个附录'):
        volume._page_bounds(PDF(), word)
    doc = Document(word)
    for name in (MAIN_BOOKMARK, BACK_BOOKMARK):
        for element in doc.element.xpath('.//w:bookmarkStart[@w:name="'+name+'"]'):
            element.getparent().remove(element)
    doc.save(word)
    with pytest.raises(volume.VolumeError, match='重新生成'):
        volume._page_bounds(PDF(), word)


@pytest.mark.parametrize('fence', ['```', '~~~~'])
def test_comments_hidden_only_outside_code_fences(tmp_path, fence):
    body = '# 1 范围\n<!-- INTERNAL-COMMENT\n# 附录 A\n-->\n正文保留。\n'
    body += fence + 'html\n<!-- LITERAL-CODE -->\n' + fence + '\n'
    source = source_file(tmp_path, body)
    word = tmp_path / 'comments.docx'
    render_document(source, word, profile='draft')
    text = '\n'.join(p.text for p in Document(word).paragraphs)
    assert 'INTERNAL-COMMENT' not in text
    assert '<!-- LITERAL-CODE -->' in text
    assert '正文保留。' in text
    assert not split_content(body).back_matter


def test_expansion_example_in_code_is_not_an_expansion_heading(tmp_path):
    doc = parse_markdown(source_file(tmp_path, '# 1 范围\n```text\n### REQ-001 正文展开 1\n```\n'))
    issues = volume.markdown_volume_issues(doc, 'SRS', 'large', 'review')
    assert not any(i['code'] == 'VOLUME_GENERATED_EXPANSION' for i in issues)
