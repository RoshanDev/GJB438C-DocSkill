from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
import json
import re
from typing import Any
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree

from .render import BOOKMARK_NAME

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}
PLACEHOLDER_RE = re.compile(r"(?:\bTODO\b|\bTBD\b|待补充|待确认|XXXX+|\{\{[^}\n]+\}\})", re.I)


@dataclass(slots=True)
class DocxIssue:
    severity: str
    code: str
    message: str

    def as_dict(self) -> dict[str, str]:
        return {"severity": self.severity, "code": self.code, "message": self.message}


@dataclass(slots=True)
class DocxAuditReport:
    path: Path
    profile: str
    issues: list[DocxIssue] = field(default_factory=list)
    metrics: dict[str, Any] = field(default_factory=dict)

    @property
    def errors(self) -> list[DocxIssue]:
        return [issue for issue in self.issues if issue.severity == "ERROR"]

    @property
    def warnings(self) -> list[DocxIssue]:
        return [issue for issue in self.issues if issue.severity == "WARN"]

    @property
    def passed(self) -> bool:
        return not self.errors

    def as_dict(self) -> dict[str, Any]:
        return {
            "path": str(self.path),
            "profile": self.profile,
            "passed": self.passed,
            "summary": {"errors": len(self.errors), "warnings": len(self.warnings)},
            "metrics": self.metrics,
            "issues": [issue.as_dict() for issue in self.issues],
        }

    def to_json(self) -> str:
        return json.dumps(self.as_dict(), ensure_ascii=False, indent=2)

    def to_text(self) -> str:
        status = "PASS" if self.passed else "FAIL"
        lines = [
            f"[{status}] {self.path} profile={self.profile}",
            f"errors={len(self.errors)} warnings={len(self.warnings)} metrics={self.metrics}",
        ]
        lines.extend(f"- {issue.severity} {issue.code}: {issue.message}" for issue in self.issues)
        return "\n".join(lines)


def _add(report: DocxAuditReport, severity: str, code: str, message: str) -> None:
    report.issues.append(DocxIssue(severity, code, message))


def _severity(profile: str) -> str:
    return "ERROR" if profile == "release" else "WARN"


def _font_info(style) -> tuple[str | None, str | None, float | None]:
    rpr = style._element.rPr
    east = None
    latin = style.font.name
    if rpr is not None and rpr.rFonts is not None:
        east = rpr.rFonts.get(qn("w:eastAsia"))
        latin = rpr.rFonts.get(qn("w:ascii")) or rpr.rFonts.get(qn("w:hAnsi")) or latin
    size = style.font.size.pt if style.font.size is not None else None
    return east, latin, size


def _line_spacing_value(style) -> float | None:
    value = style.paragraph_format.line_spacing
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _check_style(
    document: Document,
    report: DocxAuditReport,
    name: str,
    *,
    east: str,
    latin: str,
    size: float,
    line: float,
    alignment=None,
) -> None:
    try:
        style = document.styles[name]
    except KeyError:
        _add(report, "ERROR", "STYLE_MISSING", f"缺少样式 {name}")
        return
    actual_east, actual_latin, actual_size = _font_info(style)
    if actual_east != east:
        _add(report, "ERROR", "STYLE_EAST_ASIA_FONT", f"{name} 东亚字体为 {actual_east!r}，期望 {east}")
    if actual_latin != latin:
        _add(report, "ERROR", "STYLE_LATIN_FONT", f"{name} 西文字体为 {actual_latin!r}，期望 {latin}")
    if actual_size is None or abs(actual_size - size) > 0.1:
        _add(report, "ERROR", "STYLE_SIZE", f"{name} 字号为 {actual_size}pt，期望 {size}pt")
    actual_line = _line_spacing_value(style)
    if actual_line is None or abs(actual_line - line) > 0.05:
        _add(report, "ERROR", "STYLE_LINE_SPACING", f"{name} 行距为 {actual_line}，期望 {line}")
    if alignment is not None and style.paragraph_format.alignment != alignment:
        _add(report, "ERROR", "STYLE_ALIGNMENT", f"{name} 对齐方式不符合约束")


def audit_docx(path: str | Path, *, profile: str = "review") -> DocxAuditReport:
    if profile not in {"review", "release"}:
        raise ValueError("profile 必须是 review 或 release")
    source = Path(path)
    report = DocxAuditReport(source, profile)
    try:
        with ZipFile(source) as archive:
            document_xml = archive.read("word/document.xml")
            settings_xml = archive.read("word/settings.xml")
            document_rels_xml = archive.read("word/_rels/document.xml.rels")
            package_names = set(archive.namelist())
            footer_parts = {name: archive.read(name) for name in package_names if name.startswith("word/footer") and name.endswith(".xml")}
    except (FileNotFoundError, BadZipFile, KeyError) as exc:
        _add(report, "ERROR", "DOCX_PACKAGE", f"DOCX 包损坏或缺少必要部件：{exc}")
        return report

    root = etree.fromstring(document_xml)
    settings = etree.fromstring(settings_xml)
    body = root.find("./w:body", NS)
    if body is None:
        _add(report, "ERROR", "BODY_MISSING", "document.xml 缺少 body")
        return report

    text = "".join(root.xpath(".//w:t/text()", namespaces=NS))
    top_tables = [element for element in body if element.tag == f"{{{W}}}tbl"]
    if len(top_tables) < 2:
        _add(report, "ERROR", "FRONT_TABLES", "首页/签字页/变更履历结构不完整")
    else:
        if len(top_tables[0].findall("./w:tr", NS)) != 5:
            _add(report, "ERROR", "FRONT_MAIN_TABLE", "首页/签字页主表不是 5 行")
        if len(top_tables[1].findall("./w:tr", NS)) < 2:
            _add(report, "ERROR", "REVISION_TABLE", "变更履历表缺少数据行")
        first_revision_cells = top_tables[1].findall("./w:tr[1]/w:tc", NS)
        if len(first_revision_cells) != 4:
            _add(report, "ERROR", "REVISION_COLUMNS", "变更履历表不是 4 列")

    for marker in ("档案编号：", "签 署 页", "变更履历", "编制：", "批准："):
        if marker not in text:
            _add(report, "ERROR", "FRONT_MARKER", f"缺少前三页标记：{marker}")

    instructions = " ".join(root.xpath(".//w:instrText/text()", namespaces=NS))
    if "TOC" not in instructions:
        _add(report, "ERROR", "TOC_FIELD", "缺少原生 Word TOC 域")
    toc_cache_stale = "请在 Word/WPS 中更新目录" in text
    if toc_cache_stale:
        _add(
            report,
            _severity(profile),
            "TOC_CACHE_STALE",
            "TOC 域存在，但可见目录仍是待更新占位文本；发布前运行 gjb438c refresh-toc 或在 Word/WPS 中更新域",
        )

    # PAGE fields live in footer parts, not document.xml. Count only footers
    # referenced by actual section properties so stale, unreferenced template
    # parts cannot make a broken document pass.
    rel_root = etree.fromstring(document_rels_xml)
    rel_map = {rel.get("Id"): rel.get("Target") for rel in rel_root}
    referenced_footers: set[str] = set()
    relationship_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    for ref in root.xpath('.//w:sectPr/w:footerReference', namespaces=NS):
        rid = ref.get(f"{{{relationship_ns}}}id")
        target = rel_map.get(rid)
        if target:
            referenced_footers.add("word/" + target.lstrip("/"))
    page_field_count = 0
    for part in referenced_footers:
        payload = footer_parts.get(part)
        if payload:
            footer_root = etree.fromstring(payload)
            page_field_count += sum(
                1 for value in footer_root.xpath('.//w:instrText/text()', namespaces=NS)
                if "PAGE" in value.upper()
            )
    if page_field_count < 2:
        _add(report, "ERROR", "PAGE_FIELD", "目录/正文页脚缺少足够的原生 PAGE 域")
    toc_anchor_names = root.xpath(
        './/w:hyperlink/@w:anchor', namespaces=NS
    )
    bookmark_names = set(root.xpath('.//w:bookmarkStart/@w:name', namespaces=NS))
    missing_toc_targets = sorted({name for name in toc_anchor_names if name not in bookmark_names})
    if missing_toc_targets:
        _add(
            report,
            _severity(profile),
            "TOC_LINK_TARGET_MISSING",
            f"可见目录中有 {len(missing_toc_targets)} 个链接缺少正文书签目标",
        )

    bookmarks = root.xpath(
        './/w:bookmarkStart[@w:name=$name]', namespaces=NS, name=BOOKMARK_NAME
    )
    if len(bookmarks) != 1:
        _add(report, "ERROR", "BODY_BOOKMARK", f"正文书签 {BOOKMARK_NAME} 命中 {len(bookmarks)} 次")
    update = settings.find("./w:updateFields", NS)
    if update is None or update.get(f"{{{W}}}val") not in {"true", "1"}:
        _add(report, "ERROR", "UPDATE_FIELDS", "settings.xml 未启用打开时更新域")

    placeholders = PLACEHOLDER_RE.findall(text)
    if placeholders:
        _add(report, _severity(profile), "PLACEHOLDER", f"DOCX 中仍有 {len(placeholders)} 处占位内容")

    try:
        document = Document(source)
    except Exception as exc:  # pragma: no cover - python-docx internals
        _add(report, "ERROR", "DOCX_OPEN", f"python-docx 无法打开：{exc}")
        return report

    _check_style(
        document, report, "GJB 目录标题", east="宋体", latin="Times New Roman", size=16, line=1.5,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    toc_style_name = next(
        (style.name for style in document.styles if style.name.lower() == "toc 1"),
        "TOC 1",
    )
    _check_style(
        document, report, toc_style_name, east="宋体", latin="Times New Roman", size=12, line=1.5,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    _check_style(
        document, report, "GJB 正文", east="宋体", latin="Times New Roman", size=12, line=1.5,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    _check_style(
        document, report, "GJB 图表题", east="黑体", latin="Times New Roman", size=10.5, line=1.0,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _check_style(
        document, report, "GJB 表内文字", east="宋体", latin="Times New Roman", size=10.5, line=1.0,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    for level in range(1, 10):
        _check_style(
            document,
            report,
            f"GJB 标题 {level}",
            east="黑体",
            latin="Times New Roman",
            size=12,
            line=1.5,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
        )

    body_style_counts: dict[str, int] = {}
    active = False
    for paragraph in document.paragraphs:
        if paragraph._p.xpath(f'.//w:bookmarkStart[@w:name="{BOOKMARK_NAME}"]'):
            active = True
        if active:
            name = paragraph.style.name if paragraph.style is not None else ""
            if paragraph.text.strip():
                body_style_counts[name] = body_style_counts.get(name, 0) + 1
                allowed = {
                    "GJB 正文", "GJB 正文无缩进", "GJB 图表题", "GJB 代码",
                    *(f"GJB 标题 {level}" for level in range(1, 10)),
                }
                if name not in allowed:
                    _add(report, "ERROR", "BODY_STYLE", f"正文段落使用未允许样式 {name!r}: {paragraph.text[:40]}")
        if active and paragraph._p.xpath('.//w:bookmarkEnd'):
            break

    report.metrics.update(
        {
            "top_level_tables": len(top_tables),
            "toc_field": "TOC" in instructions,
            "toc_cache_stale": toc_cache_stale,
            "toc_links": len(toc_anchor_names),
            "missing_toc_link_targets": len(missing_toc_targets),
            "page_fields": page_field_count,
            "referenced_footers": sorted(referenced_footers),
            "body_style_counts": body_style_counts,
            "sections": len(document.sections),
        }
    )
    if len(document.sections) < 3:
        _add(report, "ERROR", "SECTION_COUNT", "应至少包含前三页、目录、正文三个节")
    return report
